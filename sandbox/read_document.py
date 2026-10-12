#!/usr/bin/env python3
"""Extract the text of a document. cedar's `read_document` tool runs this inside the sandbox; it is
installed in the image as `cedar-read-document`.

    cedar-read-document [--pages RANGE] [--sheet NAME] [--rows RANGE] -- PATH

Supported: PDF (pypdf), Word .docx (python-docx), Excel .xlsx/.xlsm (openpyxl), PowerPoint .pptx
(python-pptx), CSV/TSV, and HTML. The format is chosen by file content where it can be (PDF and
the zip-based Office formats), otherwise by extension. Tables are printed as Markdown tables.
RANGE is 1-based: "3", "2-5", "7-", or a comma-separated list of those. --pages selects PDF pages or
slides; --sheet and --rows select spreadsheet data (the header row is always printed).

The text goes to stdout, starting with a one-line summary of the document. Output stops after
MAX_OUTPUT characters, since cedar cuts it much shorter anyway. Warnings go to stderr; a failure
prints its reason to stderr and exits 1.
"""

import argparse
import csv
import html
import io
import os
import re
import sys
import zipfile

MAX_OUTPUT = 2_000_000


class Output:
    """Collects text and stops the program once there is more than cedar could show."""

    def __init__(self):
        self.size = 0

    def write(self, text=""):
        sys.stdout.write(text + "\n")
        self.size += len(text) + 1
        if self.size > MAX_OUTPUT:
            sys.stdout.write("\n[output stopped here; select a smaller part]\n")
            sys.stdout.flush()
            sys.exit(0)


def fail(message):
    print(message, file=sys.stderr)
    sys.exit(1)


def warn(message):
    print(message, file=sys.stderr)


def parse_spans(spec):
    """`spec` as a list of (start, end) spans, 1-based and inclusive; end is None for "N-"."""
    spans = []
    for part in spec.split(","):
        part = part.strip()
        if "-" in part:
            a, b = (x.strip() for x in part.split("-", 1))
            start, end = int(a), (int(b) if b else None)
        else:
            start = end = int(part)
        if start < 1 or (end is not None and end < start):
            raise ValueError(part)
        spans.append((start, end))
    return spans


def parse_range(spec, total):
    """The numbers up to `total` that `spec` selects, sorted, or None for "everything"."""
    if not spec:
        return None
    selected = set()
    for start, end in parse_spans(spec):
        selected.update(range(start, min(end or total, total) + 1))
    return sorted(selected)


def describe(numbers):
    """[1, 2, 3, 7] -> "1-3, 7"."""
    spans = []
    for n in numbers:
        if spans and n == spans[-1][1] + 1:
            spans[-1][1] = n
        else:
            spans.append([n, n])
    return ", ".join(str(a) if a == b else f"{a}-{b}" for a, b in spans)


def cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).replace("\r\n", "\n").replace("\n", "<br>").replace("|", "\\|")
    return text.strip()


def table(rows, out):
    """Print rows (lists of values) as a Markdown table, the first row as the header."""
    rows = [[cell(v) for v in r] for r in rows]
    if not rows:
        return
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out.write("| " + " | ".join(rows[0]) + " |")
    out.write("|" + "---|" * width)
    for r in rows[1:]:
        out.write("| " + " | ".join(r) + " |")


def plural(n, word):
    return f"{n} {word}{'' if n == 1 else 's'}"


# ----------------------------------------------------------------------------------------------


def read_pdf(path, args, out):
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                fail(f"{path} is encrypted with a password")
        except Exception as e:  # noqa: BLE001
            fail(f"{path} is encrypted and cannot be read ({e})")
    total = len(reader.pages)
    pages = parse_range(args.pages, total)
    if pages is None:
        pages = list(range(1, total + 1))
    elif not pages:
        fail(f"{path} has {plural(total, 'page')}; nothing matches pages={args.pages}")
    title = (reader.metadata or {}).get("/Title") if reader.metadata else None
    summary = f"{path}: PDF, {plural(total, 'page')}"
    if title:
        summary += f", title {str(title).strip()!r}"
    if len(pages) < total:
        summary += f" (showing {describe(pages)})"
    out.write(summary)
    empty = []
    for n in pages:
        try:
            text = reader.pages[n - 1].extract_text() or ""
        except Exception as e:  # noqa: BLE001 - a broken page shouldn't lose the rest
            warn(f"page {n}: {e}")
            text = ""
        out.write()
        out.write(f"--- page {n} ---")
        if text.strip():
            out.write(text.strip())
        else:
            empty.append(n)
            out.write("[no text]")
    if empty and len(empty) == len(pages):
        warn("no text could be extracted; the PDF is probably scanned images (no OCR is available)")
    elif empty:
        warn(f"no text on page {describe(empty)} (images or scans)")


def iter_docx_blocks(doc):
    """Paragraphs and tables in document order."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield Paragraph(child, doc)
        elif tag == "tbl":
            yield Table(child, doc)


def read_docx(path, args, out):
    import docx

    doc = docx.Document(path)
    blocks = list(iter_docx_blocks(doc))
    n_tables = sum(1 for b in blocks if b.__class__.__name__ == "Table")
    out.write(f"{path}: Word document, {plural(n_tables, 'table')}")
    out.write()
    for block in blocks:
        if block.__class__.__name__ == "Table":
            rows = [[c.text for c in row.cells] for row in block.rows]
            out.write()
            table(rows, out)
            out.write()
            continue
        text = block.text.strip()
        if not text:
            continue
        style = (block.style.name if block.style is not None else "") or ""
        m = re.match(r"Heading (\d)", style)
        if style == "Title":
            out.write(f"# {text}")
        elif m:
            out.write(f"{'#' * min(int(m.group(1)) + 1, 6)} {text}")
        elif style.startswith("List"):
            out.write(f"- {text}")
        else:
            out.write(text)
        out.write()


def sheet_rows(rows, args, out, label):
    """Print the header row plus the selected data rows from an iterator of row sequences."""
    rows = iter(rows)
    header = next(rows, None)
    if header is None:
        out.write(f"## {label} (empty)")
        out.write()
        return
    spans = parse_spans(args.rows) if args.rows else None
    # Without an open-ended span, reading can stop after the last selected row.
    last = None if spans is None or any(e is None for _, e in spans) else max(e for _, e in spans)
    picked, count, complete = [list(header)], 0, True
    for i, row in enumerate(rows, start=1):
        if last is not None and i > last:
            complete = False
            break
        count = i
        if spans is None or any(s <= i and (e is None or i <= e) for s, e in spans):
            picked.append(list(row))
    # Trailing empty columns (openpyxl reports formatted-but-empty cells) only add noise.
    width = max(
        (max((j + 1 for j, v in enumerate(r) if v not in (None, "")), default=0) for r in picked),
        default=0,
    )
    picked = [r[:width] for r in picked]
    total = plural(count, "data row") if complete else f"at least {plural(count, 'data row')}"
    if spans is None:
        out.write(f"## {label} ({total})")
    else:
        out.write(f"## {label} ({total}; showing {plural(len(picked) - 1, 'row')} of rows {args.rows})")
    out.write()
    if width:
        table(picked, out)
    out.write()


def read_xlsx(path, args, out):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    names = wb.sheetnames
    if args.sheet:
        matches = [n for n in names if n == args.sheet] or [
            n for n in names if n.lower() == args.sheet.lower()
        ]
        if not matches:
            fail(f"{path} has no sheet {args.sheet!r}; its sheets are: {', '.join(names)}")
        names = matches[:1]
    out.write(f"{path}: Excel workbook, {plural(len(wb.sheetnames), 'sheet')}: {', '.join(wb.sheetnames)}")
    out.write()
    for name in names:
        ws = wb[name]
        sheet_rows(ws.iter_rows(values_only=True), args, out, f"Sheet: {name}")
    wb.close()


def read_pptx(path, args, out):
    from pptx import Presentation

    prs = Presentation(path)
    total = len(prs.slides)
    slides = parse_range(args.pages, total)
    if slides is None:
        slides = list(range(1, total + 1))
    elif not slides:
        fail(f"{path} has {plural(total, 'slide')}; nothing matches pages={args.pages}")
    summary = f"{path}: PowerPoint, {plural(total, 'slide')}"
    if len(slides) < total:
        summary += f" (showing {describe(slides)})"
    out.write(summary)
    for n in slides:
        slide = prs.slides[n - 1]
        out.write()
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        out.write(f"--- slide {n}{': ' + title if title else ''} ---")
        title_id = title_shape.shape_id if title_shape is not None else None
        for shape in slide.shapes:
            if shape.shape_id == title_id:
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                table([[c.text for c in row.cells] for row in shape.table.rows], out)
            elif shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    text = "".join(r.text for r in p.runs).strip()
                    if text:
                        out.write(("  " * p.level) + "- " + text)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.write(f"[notes] {notes}")


def read_delimited(path, args, out, delimiter=None):
    with open(path, "rb") as f:
        raw = f.read()
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    name = {",": "CSV", "\t": "TSV", ";": "CSV (semicolon-separated)", "|": "CSV (pipe-separated)"}[delimiter]
    out.write(f"{path}: {name}")
    out.write()
    sheet_rows(csv.reader(io.StringIO(text), delimiter=delimiter), args, out, "Rows")


class _Tags:
    SKIP = {"script", "style", "noscript", "template", "svg", "head"}
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section",
             "article", "header", "footer", "blockquote", "pre", "table", "ul", "ol", "dt", "dd"}


def read_html(path, args, out):
    from html.parser import HTMLParser

    class Parser(HTMLParser):
        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.parts, self.skip, self.title, self.in_title = [], 0, "", False

        def handle_starttag(self, tag, attrs):
            if tag in _Tags.SKIP:
                self.skip += 1
            if tag == "title":
                self.in_title = True
            if tag in _Tags.BLOCK:
                self.parts.append("\n")
            m = re.fullmatch(r"h([1-6])", tag)
            if m and not self.skip:
                self.parts.append("#" * int(m.group(1)) + " ")
            if tag == "li" and not self.skip:
                self.parts.append("- ")

        def handle_endtag(self, tag):
            if tag in _Tags.SKIP and self.skip:
                self.skip -= 1
            if tag == "title":
                self.in_title = False
            if tag in _Tags.BLOCK:
                self.parts.append("\n")

        def handle_data(self, data):
            if self.in_title:
                self.title += data
            elif not self.skip:
                self.parts.append(data)

    with open(path, "rb") as f:
        source = f.read().decode("utf-8", errors="replace")
    p = Parser()
    p.feed(source)
    text = "".join(p.parts)
    lines = [re.sub(r"[ \t\r\f\v]+", " ", l).strip() for l in text.split("\n")]
    text = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
    summary = f"{path}: HTML"
    if p.title.strip():
        summary += f", title {html.unescape(p.title.strip())!r}"
    out.write(summary)
    out.write()
    out.write(text)


# ----------------------------------------------------------------------------------------------

READERS = {
    "pdf": read_pdf,
    "docx": read_docx,
    "xlsx": read_xlsx,
    "pptx": read_pptx,
    "csv": read_delimited,
    "tsv": lambda p, a, o: read_delimited(p, a, o, "\t"),
    "html": read_html,
}

EXTENSIONS = {
    ".pdf": "pdf", ".docx": "docx", ".docm": "docx", ".xlsx": "xlsx", ".xlsm": "xlsx",
    ".pptx": "pptx", ".pptm": "pptx", ".csv": "csv", ".tsv": "tsv", ".tab": "tsv",
    ".html": "html", ".htm": "html", ".xhtml": "html",
}

LEGACY = {".doc": "Word 97-2003", ".xls": "Excel 97-2003", ".ppt": "PowerPoint 97-2003"}


def detect(path):
    """The reader for `path`: by magic bytes for PDF and Office zips, otherwise by extension."""
    with open(path, "rb") as f:
        head = f.read(1024)
    if head.startswith(b"%PDF-"):
        return "pdf"
    if head.startswith(b"PK") and zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as z:
            names = set(z.namelist())
        if "word/document.xml" in names:
            return "docx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        if "ppt/presentation.xml" in names:
            return "pptx"
    ext = os.path.splitext(path)[1].lower()
    if ext in LEGACY:
        fail(f"{path} is a legacy {LEGACY[ext]} file, which can't be read here; save it as "
             f"{ext}x instead")
    if ext in EXTENSIONS:
        return EXTENSIONS[ext]
    if re.match(rb"\s*(<!doctype html|<html)", head, re.I):
        return "html"
    fail(f"{path}: unsupported format; read_document handles PDF, .docx, .xlsx, .pptx, CSV/TSV, "
         "and HTML (use read_file for plain text)")


def main():
    parser = argparse.ArgumentParser(description="Extract the text of a document.")
    parser.add_argument("--pages")
    parser.add_argument("--sheet")
    parser.add_argument("--rows")
    parser.add_argument("path")
    args = parser.parse_args()

    for name in ("pages", "rows"):
        value = getattr(args, name)
        if value:
            try:
                parse_spans(value)
            except ValueError:
                fail(f"invalid {name} {value!r}; use forms like 3, 2-5, 7-, or 1,4-6")

    if not os.path.exists(args.path):
        fail(f"{args.path}: no such file")
    if os.path.isdir(args.path):
        fail(f"{args.path} is a directory")

    kind = detect(args.path)
    out = Output()
    try:
        READERS[kind](args.path, args, out)
    except ImportError as e:  # pragma: no cover - depends on the image
        fail(f"reading {kind} files is not available in this sandbox image ({e})")
    except SystemExit:
        raise
    except Exception as e:  # noqa: BLE001 - whatever the parser raises is the reason to report
        fail(f"cannot read {args.path} as {kind}: {e}")


if __name__ == "__main__":
    main()
